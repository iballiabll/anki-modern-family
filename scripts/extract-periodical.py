"""把 D:\\外刊 里的 PDF / DOCX 学习材料解析成结构化 JSON。

用法：
    python scripts/extract-periodical.py --root "D:\\外刊" --out ".cache/periodical-raw"

产物是「每个源文件一个 JSON」，由 scripts/build-periodical-library.mjs 合并成
periodical-data/*.js。解析结果只做文本级重排，不还原杂志版式，遇到无法解析的
扫描件会写入 warnings，交由构建脚本汇总。
"""

from __future__ import annotations

import argparse
import json
import pathlib
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

CJK_RE = re.compile(r"[\u3400-\u9fff\u3000-\u303f\uff00-\uffef]")
EN_FONT_HINTS = ("times", "arial", "cambria", "georgia", "calibri", "garamond", "helvetica")
# 版式里段落标记常见多种写法：Para.1 / 【Para.1】 / 【️【️Para.1】
PARA_PREFIX = r"^[\s【\ufe0f]*"
PARA_MARK_RE = re.compile(PARA_PREFIX + r"[Pp]ara\.?\s*([0-9]{1,2})\s*】?\s*$")
PARA_INLINE_RE = re.compile(PARA_PREFIX + r"[Pp]ara\.?\s*([0-9]{1,2})\s*】\s*(.*)$")
VOCAB_MARK = "⚫"
SECTION_MARKS = ("➢", "【️", "【")
DATE_IN_NAME_RE = re.compile(r"^(?:【福利】)?\s*(\d{1,2})\.(\d{1,2})")
RANGE_IN_NAME_RE = re.compile(r"(\d{1,2})\.(\d{1,2})\s*[—\-~至]\s*(\d{1,2})\.(\d{1,2})")
SOURCE_HEAD_RE = re.compile(r"^【\s*(\d{1,2}\.\d{1,2})\s*\|\s*([^】]+)】\s*(.*)$")
ARTICLE_FOOT_RE = re.compile(r"^(.+?)\s*\|\s*([A-Za-z][A-Za-z .'&-]{2,40})$")
OPTION_LINE_RE = re.compile(r"\[([A-G])\]\s*(.+)$")
BLANK_OPTION_RE = re.compile(r"\[([A-D])\]\s*([^\[\]]+?)(?=\s*\[[A-D]\]|$)")
# 不同期次的检验题把选项写成 (A) / A. / A、 / A） 等，先统一成 [A] 再走同一套解析。
BRACKET_OPTION_RE = re.compile(r"[（(]\s*([A-G])\s*[)）]")
LINE_OPTION_RE = re.compile(r"^([A-G])\s*[.、)）]\s*")
QUESTION_RE = re.compile(r"^(\d{1,2})[.、]\s*(.+)$")
# 检验题 DOCX 的小节标题既有中文写法，也有英文写法（Answer Key / Detailed Explanations）。
ANSWER_SECTION_MARKS = ("答案速查", "答案与解析速查", "Answer Key", "Answer key", "Answers", "Keys")
ANALYSIS_SECTION_MARKS = (
    "详细解析",
    "详细答案解析",
    "文章与命题分析",
    "Detailed Explanations",
    "Detailed Explanation",
    "Detailed Analysis",
    "Explanations",
)
SKIP_SECTION_MARKS = ("答案与解析", "Multiple Choice", "Answer Sheet")
BOILERPLATE = ("版权所有", "禁止传播", "公众号：", "关注微信公众号", "扫码", "第 ", "微信")
PROMO_INLINE = (
    "关注微信公众号",
    "公众号【",
    "加入训练营",
    "扫码",
    "资料禁言群",
    "本资料",
    "禁止传播",
    "版权所有",
)
PROMO_LINE = (
    "本文档由我们自研",
    "本文档为【福利】",
    "训练营创办",
    "如果对答案解析有疑惑",
    "欢迎大家自行交流使用",
    "该模型基于近10年考研英语真题",
    "我们会根据同学们的反馈持续优化",
)


def log(message: str) -> None:
    print(message, flush=True)


def is_cjk(text: str) -> bool:
    return bool(CJK_RE.search(text))


def dominant_font(chars) -> str:
    counter: dict[str, int] = {}
    for char in chars:
        name = (char.get("fontname") or "").lower()
        counter[name] = counter.get(name, 0) + 1
    if not counter:
        return ""
    return max(counter.items(), key=lambda item: item[1])[0]


def line_kind(text: str, font: str) -> str:
    if is_cjk(text):
        return "zh"
    if any(hint in font for hint in EN_FONT_HINTS):
        return "en"
    return "en" if re.search(r"[A-Za-z]{2}", text) else "zh"


def rebuild_line(chars) -> str:
    """按字符间距重建文本，修复部分 PDF 丢失空格的问题。"""
    out: list[str] = []
    previous = None
    for char in chars:
        if previous is not None:
            gap = char["x0"] - previous["x1"]
            size = max(float(char.get("size") or 0), float(previous.get("size") or 0), 1.0)
            text = char.get("text") or ""
            prev_text = previous.get("text") or ""
            if (
                gap > size * 0.14
                and gap > 0.6
                and prev_text
                and text
                and not prev_text[-1].isspace()
                and not text[0].isspace()
                and not (CJK_RE.search(prev_text[-1]) and CJK_RE.search(text[0]))
            ):
                out.append(" ")
        out.append(char.get("text") or "")
        previous = char
    return "".join(out).strip()


def detect_column_split(page) -> float | None:
    """判断页面是否为双栏排版，是则返回分栏位置。"""
    width = float(page.width)
    middle = width * 0.5
    gutters = 0
    total = 0
    for raw in page.extract_text_lines(return_chars=True):
        chars = raw.get("chars") or []
        if len(chars) < 8:
            continue
        total += 1
        left = [char for char in chars if float(char["x0"]) < middle]
        right = [char for char in chars if float(char["x0"]) >= middle]
        if not left or not right:
            continue
        if min(float(char["x0"]) for char in right) - max(float(char["x1"]) for char in left) > 5:
            gutters += 1
    if total >= 3 and gutters / total > 0.5:
        return width * 0.53
    return None


def chunk_text(chunk) -> str:
    plain = "".join(char.get("text") or "" for char in chunk).strip()
    if not plain:
        return ""
    length = max(len(plain), 1)
    cjk_ratio = sum(1 for char in plain if CJK_RE.search(char)) / length
    space_ratio = plain.count(" ") / length
    if cjk_ratio > 0.3 or space_ratio >= 0.06:
        return plain
    return rebuild_line(chunk)


def pdf_lines(path: pathlib.Path) -> tuple[list[dict], int, str]:
    import pdfplumber

    lines: list[dict] = []
    with pdfplumber.open(str(path)) as pdf:
        page_count = len(pdf.pages)
        first_text = ""
        for page_index, page in enumerate(pdf.pages, start=1):
            split = detect_column_split(page)
            raw_lines = page.extract_text_lines(return_chars=True)
            for raw in raw_lines:
                chars = raw.get("chars") or []
                if not chars:
                    continue
                if split is None:
                    buckets = {0: chars}
                else:
                    buckets = {}
                    for char in chars:
                        key = 0 if float(char["x0"]) < split else 1
                        buckets.setdefault(key, []).append(char)
                for key in sorted(buckets):
                    chunk = sorted(buckets[key], key=lambda item: float(item["x0"]))
                    text = chunk_text(chunk)
                    if not text:
                        continue
                    font = dominant_font(chunk)
                    kind = line_kind(text, font)
                    lines.append(
                        {
                            "page": page_index,
                            "column": key,
                            "top": round(float(raw["top"]), 1),
                            "kind": kind,
                            "font": font,
                            "text": text,
                        }
                    )
                    if not first_text and len(text) > 40:
                        first_text = text
    return lines, page_count, first_text


def docx_paragraphs(path: pathlib.Path) -> tuple[list[str], list[list[list[str]]]]:
    import docx

    document = docx.Document(str(path))
    paragraphs = [(paragraph.text or "").strip() for paragraph in document.paragraphs]
    tables: list[list[list[str]]] = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append(rows)
    return paragraphs, tables


def is_boilerplate(text: str) -> bool:
    if re.fullmatch(r"[0-9]{1,4}", text):
        return True
    if any(text.startswith(mark) for mark in BOILERPLATE) or "版权所有" in text:
        return True
    return any(mark in text for mark in PROMO_LINE)


def strip_promo(text: str) -> str:
    """去掉正文/解析尾部拼接的公众号推广语。"""
    cut = len(text)
    for mark in PROMO_INLINE:
        at = text.find(mark)
        if at >= 0:
            cut = min(cut, at)
    cleaned = text[:cut].strip()
    # 推广语可能被切成多个段落，尾部只留下“关注微信”之类的残片。
    if len(cleaned) > 60:
        tail = re.search(r"(关注|扫码|微信|公众号|训练营|禁言群|版权所有|禁止传播)[^。！？!?]{0,40}$", cleaned)
        if tail:
            cleaned = cleaned[: tail.start()].strip()
    return cleaned


def classify(name: str) -> dict:
    stem = pathlib.Path(name).stem
    info = {
        "name": name,
        "stem": stem,
        "ext": pathlib.Path(name).suffix.lower(),
        "kind": "",
        "testType": "",
        "date": "",
        "range": "",
    }
    if "杂志排版" in stem or re.search(r"排版", stem):
        info["kind"] = "layout"
    if "精读" in stem:
        info["kind"] = "reading"
    if "原文" in stem:
        info["kind"] = "source"
    if "答疑" in stem:
        info["kind"] = "qa"
    if "检验题" in stem:
        info["kind"] = "test"
        if "阅读理解" in stem:
            info["testType"] = "reading"
        elif "完型" in stem or "完形" in stem:
            info["testType"] = "cloze"
        elif "7选5" in stem:
            info["testType"] = "gap-fill"
        elif "段落排序" in stem:
            info["testType"] = "ordering"
        elif "小标题" in stem:
            info["testType"] = "heading"
        else:
            info["testType"] = "other"
    match = DATE_IN_NAME_RE.match(stem)
    if match:
        info["date"] = f"2026-{int(match.group(1)):02d}-{int(match.group(2)):02d}"
    range_match = RANGE_IN_NAME_RE.search(stem)
    if range_match:
        info["range"] = (
            f"2026-{int(range_match.group(1)):02d}-{int(range_match.group(2)):02d}|"
            f"2026-{int(range_match.group(3)):02d}-{int(range_match.group(4)):02d}"
        )
        if not info["date"]:
            info["date"] = info["range"].split("|")[0]
    return info


def new_vocab(term: str, head: str) -> dict:
    return {
        "term": term,
        "head": head,
        "pos": "",
        "phonetic": "",
        "definition": [],
        "gloss": [],
        "example": [],
        "exampleZh": [],
        "tags": [],
        "phase": "definition",
    }


def finish_vocab(entry: dict) -> dict:
    entry.pop("phase", None)
    entry["definition"] = " ".join(part for part in entry["definition"] if part).strip()
    entry["gloss"] = " ".join(part for part in entry["gloss"] if part).strip()
    entry["example"] = " ".join(part for part in entry["example"] if part).strip()
    entry["exampleZh"] = " ".join(part for part in entry["exampleZh"] if part).strip()
    if not entry["example"] and entry["gloss"]:
        entry["gloss"], entry["example"] = "", entry["gloss"]
    return entry


def split_vocab_head(text: str) -> tuple[str, str]:
    cleaned = text.replace(VOCAB_MARK, "").strip()
    match = re.match(r"^([A-Za-z][A-Za-z\-' ]{0,40}?)\s+((?:n|v|vt|vi|adj|adv|prep|phr|conj|pron|num)\.\s*.*)$", cleaned)
    if match:
        return match.group(1).strip(), match.group(2).strip()
    return cleaned, ""


SECTION_LABELS = (
    "今日句子分析",
    "句子分析",
    "写作积累",
    "翻译作业",
    "参考译文",
    "拓展表达",
    "今日文章脉络",
    "长难句",
    "语法",
)


def parse_reading_stream(
    lines: list[dict], entry: dict, collect_vocab: bool = True, collect_sections: bool = True
) -> dict:
    paragraphs: dict[int, dict] = {}
    vocab: dict[str, dict] = {}
    sections: list[dict] = []
    current_para: int | None = None
    current_vocab: dict | None = None
    current_section: dict | None = None
    order: list[int] = []

    def ensure_para(number: int) -> dict:
        if number not in paragraphs:
            paragraphs[number] = {"index": number, "en": [], "zh": []}
            order.append(number)
        return paragraphs[number]

    for line in lines:
        text = line["text"].strip()
        if not text or is_boilerplate(text):
            continue
        kind = line["kind"]

        inline = PARA_INLINE_RE.match(text)
        marker = PARA_MARK_RE.match(text)
        if inline:
            current_para = int(inline.group(1))
            current_vocab = None
            current_section = None
            ensure_para(current_para)
            rest = inline.group(2).strip()
            if rest:
                target = ensure_para(current_para)
                key = "zh" if kind == "zh" else "en"
                target[key].append(rest)
            continue
        if marker:
            current_para = int(marker.group(1))
            current_vocab = None
            current_section = None
            ensure_para(current_para)
            continue

        if collect_vocab and (text.startswith(VOCAB_MARK) or re.match(
            r"^[A-Za-z][A-Za-z\-' ]{0,40}\s+(?:n|v|vt|vi|adj|adv|prep|phr|conj|pron)\.\s",
            text,
        )):
            term, head = split_vocab_head(text)
            key = term.lower()
            entry_vocab = vocab.get(key)
            if entry_vocab is None:
                entry_vocab = new_vocab(term, head)
                vocab[key] = entry_vocab
            elif head and not entry_vocab["head"]:
                entry_vocab["head"] = head
            current_vocab = entry_vocab
            current_para = None
            current_section = None
            continue

        heading = ""
        if text.startswith("➢"):
            heading = text.lstrip("➢ ").strip()
        else:
            label_match = re.match(r"^【[️\s]*([^】]{2,20})】", text)
            if label_match and any(word in label_match.group(1) for word in SECTION_LABELS):
                heading = label_match.group(1).strip()
        if heading:
            if collect_sections:
                current_section = {"heading": heading, "lines": []}
                sections.append(current_section)
            else:
                current_section = None
            current_para = None
            current_vocab = None
            continue

        if current_section is not None:
            current_section["lines"].append({"kind": kind, "text": text})
            continue

        if current_vocab is not None:
            target = current_vocab
            if kind == "zh":
                if target["phase"] == "definition":
                    target["gloss"].append(text)
                    target["phase"] = "example"
                else:
                    target["exampleZh"].append(text)
            else:
                if target["phase"] == "definition":
                    target["definition"].append(text)
                else:
                    target["example"].append(text)
            for tag in ("考研大纲词汇", "六级", "四级"):
                if tag in text and tag not in target["tags"]:
                    target["tags"].append(tag)
            continue

        if current_para is not None:
            target = paragraphs[current_para]
            if kind == "zh":
                if not target["en"] and not target["zh"]:
                    continue
                target["zh"].append(text)
            else:
                if target["zh"]:
                    # 段落译文之后又出现英文，多半是下一段的正文，另起一段。
                    current_para += 1
                    target = ensure_para(current_para)
                target["en"].append(text)
            continue

    result_paragraphs = []
    for number in sorted(paragraphs):
        block = paragraphs[number]
        en = strip_promo(" ".join(block["en"]).strip())
        zh = strip_promo(" ".join(block["zh"]).strip())
        if not en and not zh:
            continue
        result_paragraphs.append(
            {
                "index": number,
                "en": re.sub(r"\s+", " ", en).strip(),
                "zh": re.sub(r"\s+", " ", zh).strip(),
            }
        )

    result_vocab = [finish_vocab(value) for value in vocab.values() if value["term"]]
    return {
        "paragraphs": result_paragraphs,
        "vocab": result_vocab,
        "sections": [section for section in sections if section["lines"]],
        "title": entry.get("title", ""),
        "titleZh": entry.get("titleZh", ""),
        "meta": entry.get("meta", {}),
    }


def find_headline(lines: list[dict]) -> dict:
    title = ""
    title_zh = ""
    source = ""
    series = ""
    for line in lines[:120]:
        text = line["text"].strip()
        match = SOURCE_HEAD_RE.match(text)
        if match and not title:
            source = match.group(2).strip()
            title = match.group(3).strip()
            cjk_index = next(
                (index for index, char in enumerate(title) if CJK_RE.search(char)),
                -1,
            )
            if cjk_index > 3:
                trailing = title[cjk_index:]
                title = title[:cjk_index].strip()
                title_zh = re.sub(r"^[0-9\s]+", "", trailing).strip()
            continue
        if text.startswith("【") and text.endswith("】") and "系列" in text and not series:
            series = text.strip("【】")
            continue
        if is_cjk(text) and 2 <= len(text) <= 40 and not title and not title_zh and not is_boilerplate(text):
            if any(word in text for word in ("难度", "脉络", "学习流程", "考研那些事", "第一遍")):
                continue
            title_zh = text.strip("【】")
    if title:
        title = re.sub(r"\s*[0-9]{1,2}\s*$", "", title).strip()
    return {"title": title, "titleZh": title_zh, "source": source, "series": series}


def extract_pdf(path: pathlib.Path, info: dict) -> dict:
    lines, page_count, first_text = pdf_lines(path)
    payload: dict = {
        "file": info["name"],
        "kind": info["kind"],
        "testType": info["testType"],
        "date": info["date"],
        "range": info["range"],
        "pages": page_count,
        "chars": sum(len(line["text"]) for line in lines),
        "warnings": [],
    }
    if payload["chars"] < 400:
        payload["warnings"].append("文本层为空或过短，可能是扫描件")
    headline = find_headline(lines)
    payload["headline"] = headline
    if info["kind"] in {"reading", "layout"}:
        payload["reading"] = parse_reading_stream(lines, headline)
    else:
        payload["lines"] = [{"kind": line["kind"], "text": line["text"]} for line in lines]
    return payload


def normalize_option_marks(text: str) -> str:
    """把 (A) / A. / A、 / A） 这类选项标记统一成 [A]，便于同一套正则解析。"""

    if not text:
        return text
    text = BRACKET_OPTION_RE.sub(lambda match: f"[{match.group(1)}] ", text)
    return LINE_OPTION_RE.sub(lambda match: f"[{match.group(1)}] ", text)


def parse_test_docx(paragraphs: list[str], info: dict) -> dict:
    questions: dict[int, dict] = {}
    answer_key: dict[int, str] = {}
    blocks: list[dict] = []
    analysis: dict[int, dict] = {}
    passage: list[str] = []
    current: dict | None = None
    section = "passage"
    expected = 21 if info["testType"] == "reading" else 41 if info["testType"] in {"gap-fill", "ordering", "heading"} else 11

    def flush() -> None:
        nonlocal current
        if current and current["stem"]:
            questions[current["number"]] = current
        current = None

    for raw in paragraphs:
        text = normalize_option_marks(raw.strip())
        if not text:
            continue
        if text.startswith(ANSWER_SECTION_MARKS):
            flush()
            section = "answers"
            continue
        if text.startswith(ANALYSIS_SECTION_MARKS):
            flush()
            section = "analysis"
            continue
        if text.startswith(SKIP_SECTION_MARKS):
            continue
        if section == "answers":
            for number, letter in re.findall(r"(\d{1,2})\.?\s*([A-Ga-g])(?![A-Za-z])", text):
                answer_key[int(number)] = letter.upper()
            continue

        if section == "passage":
            if info["testType"] == "cloze":
                options = BLANK_OPTION_RE.findall(text)
                head = QUESTION_RE.match(text)
                if options and head:
                    questions[int(head.group(1))] = {
                        "number": int(head.group(1)),
                        "stem": "",
                        "options": [{"key": key, "text": value.strip()} for key, value in options],
                    }
                    continue
            head = QUESTION_RE.match(text)
            if head and int(head.group(1)) >= expected:
                flush()
                current = {"number": int(head.group(1)), "stem": head.group(2).strip(), "options": []}
                continue
            if current is not None:
                option_parts = BLANK_OPTION_RE.findall(text)
                if option_parts:
                    for key, value in option_parts:
                        current["options"].append({"key": key, "text": value.strip()})
                else:
                    current["stem"] = f"{current['stem']} {text}".strip()
                continue
            passage.append(text)
            continue

        if section == "analysis":
            head = re.match(r"^(\d{1,2})\.\s*(.*)$", text)
            if head:
                number = int(head.group(1))
                analysis[number] = {"kicker": head.group(2).strip(), "lines": []}
                continue
            marker = re.match(r"^【(.+?)】(.*)$", text)
            if marker and analysis:
                target = analysis[max(analysis)]
                label = marker.group(1)
                rest = marker.group(2).strip()
                if label.startswith("答案") and rest:
                    answer_key[max(analysis)] = rest.strip("[] ")
                elif rest:
                    target["lines"].append(f"【{label}】{rest}")
                else:
                    target["lines"].append(f"【{label}】")
                continue
            if analysis:
                cleaned = strip_promo(text)
                if cleaned:
                    analysis[max(analysis)]["lines"].append(cleaned)
            else:
                blocks.append({"type": "note", "text": text})
            continue

        head = QUESTION_RE.match(text)
        if head and (int(head.group(1)) >= expected or int(head.group(1)) in answer_key):
            flush()
            current = {"number": int(head.group(1)), "stem": head.group(2).strip(), "options": []}
            continue
        if current is not None:
            option_parts = BLANK_OPTION_RE.findall(text)
            if option_parts:
                for key, value in option_parts:
                    current["options"].append({"key": key, "text": value.strip()})
                continue
            current["stem"] = f"{current['stem']} {text}".strip()
            continue
        passage.append(text)
    flush()

    items = []
    for number in sorted(set(questions) | set(answer_key)):
        question = questions.get(number, {"number": number, "stem": "", "options": []})
        detail = analysis.get(number, {})
        items.append(
            {
                "number": number,
                "stem": re.sub(r"\s+", " ", question.get("stem", "")).strip(),
                "options": question.get("options", []),
                "answer": answer_key.get(number, ""),
                "kicker": detail.get("kicker", ""),
                "analysis": detail.get("lines", []),
            }
        )
    instructions = [text for text in passage[:2] if text.lower().startswith(("directions", "in the following", "read the following", "the following"))]
    return {
        "instructions": instructions,
        "passage": [text for text in passage if text not in instructions][:60],
        "items": items,
        "blocks": blocks,
    }


def extract_docx(path: pathlib.Path, info: dict) -> dict:
    paragraphs, tables = docx_paragraphs(path)
    payload: dict = {
        "file": info["name"],
        "kind": info["kind"],
        "testType": info["testType"],
        "date": info["date"],
        "range": info["range"],
        "paragraphCount": len([text for text in paragraphs if text]),
        "warnings": [],
    }
    if info["kind"] == "test":
        payload["test"] = parse_test_docx(paragraphs, info)
        payload["tables"] = [[[cell for cell in row] for row in table] for table in tables]
        return payload
    if info["kind"] == "source":
        articles: list[dict] = []
        current: dict | None = None
        for raw in paragraphs:
            text = raw.strip()
            if not text:
                continue
            head = SOURCE_HEAD_RE.match(text)
            if head:
                current = {
                    "publishedOn": head.group(1),
                    "source": head.group(2).strip(),
                    "title": head.group(3).strip(),
                    "paragraphs": [],
                    "footer": [],
                }
                articles.append(current)
                continue
            if current is None:
                continue
            marker = PARA_MARK_RE.match(text) or PARA_INLINE_RE.match(text)
            if marker:
                current["paragraphs"].append({"index": int(marker.group(1)), "text": ""})
                inline = PARA_INLINE_RE.match(text)
                if inline and inline.group(2).strip():
                    current["paragraphs"][-1]["text"] = inline.group(2).strip()
                continue
            if "正文字数" in text or ARTICLE_FOOT_RE.match(text):
                current["footer"].append(text)
                continue
            if current["paragraphs"]:
                merged = f"{current['paragraphs'][-1]['text']} {strip_promo(text)}".strip()
                current["paragraphs"][-1]["text"] = merged
        for article in articles:
            for block in article["paragraphs"]:
                block["text"] = strip_promo(block["text"])
            article["paragraphs"] = [block for block in article["paragraphs"] if block["text"]]
        payload["articles"] = articles
        payload["paragraphs"] = [text for text in paragraphs if text]
        return payload
    # 答疑汇总
    items: list[dict] = []
    current_item: dict | None = None
    mode = ""
    for raw in paragraphs:
        text = raw.strip()
        if not text:
            continue
        if re.match(r"^原文(?:句子|段落|句)?\s*[:：]", text):
            if current_item:
                items.append(current_item)
            current_item = {
                "locator": "",
                "sentence": text.split("：", 1)[-1].strip(),
                "question": "",
                "answer": "",
            }
            mode = "sentence"
            continue
        if current_item is None:
            continue
        if re.match(r"^(提问|问题|疑问)\s*[:：]?", text):
            mode = "question"
            continue
        if re.match(r"^(回答|解答|答)\s*[:：]?", text):
            mode = "answer"
            continue
        if re.match(r"^\d{1,2}\.\d{1,2}[，,]", text) and not current_item["locator"]:
            current_item["locator"] = text
            continue
        current_item[mode] = f"{current_item[mode]} {text}".strip()
    if current_item:
        items.append(current_item)
    payload["items"] = items
    payload["paragraphs"] = [text for text in paragraphs if text]
    return payload


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", default=r"D:\外刊")
    parser.add_argument("--out", default=".cache/periodical-raw")
    parser.add_argument("--limit", type=int, default=0)
    parser.add_argument("--only", action="append", default=[], help="只解析文件名包含该子串的文件")
    args = parser.parse_args()

    root = pathlib.Path(args.root)
    out_dir = pathlib.Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    files = sorted(
        [path for path in root.glob("*.pdf")] + [path for path in root.glob("*.docx")],
        key=lambda path: path.name,
    )
    if args.limit:
        files = files[: args.limit]
    if args.only:
        files = [path for path in files if any(key in path.name for key in args.only)]
    log(f"扫描到 {len(files)} 个文件，来源目录 {root}")
    manifest = []
    for index, path in enumerate(files, start=1):
        info = classify(path.name)
        try:
            payload = (
                extract_pdf(path, info)
                if info["ext"] == ".pdf"
                else extract_docx(path, info)
            )
        except Exception as error:  # noqa: BLE001 - 单个文件失败不应中断整批构建
            payload = {
                "file": path.name,
                "kind": info["kind"],
                "testType": info["testType"],
                "date": info["date"],
                "range": info["range"],
                "warnings": [f"解析失败：{type(error).__name__}: {error}"],
            }
        payload["sourcePath"] = str(path)
        payload["bytes"] = path.stat().st_size
        target = out_dir / f"{index:03d}.json"
        target.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        manifest.append({"index": index, "file": path.name, "json": target.name, "kind": info["kind"]})
        if index % 10 == 0 or index == len(files):
            log(f"  已解析 {index}/{len(files)}")
    (out_dir / "manifest.json").write_text(
        json.dumps({"root": str(root), "files": manifest}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    log(f"完成，原始结构化结果写入 {out_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
